#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
混合智能体（Hybrid Agent）- 教培机构咨询规划助手

基于LangGraph实现的混合型智能体，结合反应式架构的即时响应能力和深思熟虑架构的学习规划能力，
通过协调层动态切换处理模式，提供课程咨询与学习规划服务。

三层架构：
1. 底层（反应式）：即时响应课程、价格、班次、进度等查询
2. 中层（协调）：评估问题类型和优先级，动态选择处理模式
3. 顶层（深思熟虑）：结合学员画像做差距分析，并给出学习与续费规划

Langfuse 集成：
本智能体集成本地/自建 Langfuse 用于调试、追踪和监控。

1. 启动本地 Langfuse（Docker Compose）后访问 http://localhost:3000
2. 注册账号，创建组织与项目，在项目 Settings -> API Keys 中创建密钥

3. 设置环境变量（Windows PowerShell）：
   $env:LANGFUSE_PUBLIC_KEY="pk-lf-..."
   $env:LANGFUSE_SECRET_KEY="sk-lf-..."
   $env:LANGFUSE_BASE_URL="http://localhost:3000"

启用后，所有 LLM 调用和 Agent 执行流程都会通过 Callback 记录到 Langfuse，
可在 http://localhost:3000 查看详细的执行追踪、性能指标和调试信息。
"""

import os
import re
import json
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, Literal, TypedDict, Optional, Tuple

import openpyxl

# 修复 langchain 1.1.x 版本兼容性问题：在导入 langchain_core 之前设置必要的属性
import langchain
if not hasattr(langchain, 'verbose'):
    langchain.verbose = False
if not hasattr(langchain, 'debug'):
    langchain.debug = False
if not hasattr(langchain, 'llm_cache'):
    langchain.llm_cache = None

from langchain_core.prompts import ChatPromptTemplate
from langchain_community.llms import Tongyi
from langchain_core.output_parsers import StrOutputParser, JsonOutputParser
from langgraph.graph import StateGraph, END
from langfuse import Langfuse
from langfuse.langchain import CallbackHandler
import warnings
warnings.filterwarnings("ignore")

# 设置API密钥
DASHSCOPE_API_KEY = os.getenv("DASHSCOPE_API_KEY")

# ==================== Langfuse 配置 ====================
# Langfuse 用于 Agent 调试、追踪和监控
# 使用前需要设置以下环境变量：
# 1. LANGFUSE_PUBLIC_KEY: 项目公钥 pk-lf-...
# 2. LANGFUSE_SECRET_KEY: 项目私钥 sk-lf-...
# 3. LANGFUSE_BASE_URL: 本地一般为 http://localhost:3000
#
# 本脚本关闭 LangSmith 自动追踪，避免重复上报
os.environ["LANGCHAIN_TRACING_V2"] = "false"

LANGFUSE_PUBLIC_KEY = (os.getenv("LANGFUSE_PUBLIC_KEY") or "").strip()
LANGFUSE_SECRET_KEY = (os.getenv("LANGFUSE_SECRET_KEY") or "").strip()
LANGFUSE_BASE_URL = (os.getenv("LANGFUSE_BASE_URL") or "http://localhost:3000").strip().rstrip("/")
LANGFUSE_ENABLED = bool(LANGFUSE_PUBLIC_KEY and LANGFUSE_SECRET_KEY)

# 必须写回环境变量：CallbackHandler / OTEL 导出器读的是 os.environ，
# 仅用 Python 变量默认值时，SDK 可能仍指向云端 cloud.langfuse.com，导致本地密钥 401
if LANGFUSE_PUBLIC_KEY:
    os.environ["LANGFUSE_PUBLIC_KEY"] = LANGFUSE_PUBLIC_KEY
if LANGFUSE_SECRET_KEY:
    os.environ["LANGFUSE_SECRET_KEY"] = LANGFUSE_SECRET_KEY
os.environ["LANGFUSE_BASE_URL"] = LANGFUSE_BASE_URL
os.environ["LANGFUSE_HOST"] = LANGFUSE_BASE_URL  # 兼容旧变量名

print("LANGFUSE_ENABLED=", LANGFUSE_ENABLED)
print("LANGFUSE_BASE_URL=", LANGFUSE_BASE_URL)

# 初始化 Langfuse 客户端（启用时，显式传入 base_url，避免打到云端）
langfuse_client = None
if LANGFUSE_ENABLED:
    langfuse_client = Langfuse(
        public_key=LANGFUSE_PUBLIC_KEY,
        secret_key=LANGFUSE_SECRET_KEY,
        base_url=LANGFUSE_BASE_URL,
    )

# 创建LLM实例
llm = Tongyi(model_name="qwen-turbo-latest", dashscope_api_key=DASHSCOPE_API_KEY)

# 定义状态类型
class EducationAdvisorState(TypedDict):
    """教培咨询规划智能体的状态"""
    # 输入
    user_query: str  # 用户查询
    student_profile: Optional[Dict[str, Any]]  # 学员画像

    # 处理状态
    query_type: Optional[Literal["emergency", "informational", "analytical"]]  # 查询类型
    processing_mode: Optional[Literal["reactive", "deliberative"]]  # 处理模式
    course_data: Optional[Dict[str, Any]]  # 收集到的课程与学情数据
    analysis_results: Optional[Dict[str, Any]]  # 分析结果

    # 输出
    final_response: Optional[str]  # 最终响应

    # 控制流
    current_phase: Optional[str]
    error: Optional[str]  # 错误信息

# 提示模板
ASSESSMENT_PROMPT = """你是一个教培机构咨询规划助手的协调层。请评估以下用户查询，确定其类型和应该采用的处理模式。

用户查询: {user_query}

请判断:
1. 查询类型:
   - "emergency": 需要立即回答的事实查询（如课程价格、开课时间、班次、剩余名额、老师、已报课程进度）
   - "informational": 信息性查询，需要课程知识但不需要个人规划（如课程介绍、适合人群、两门课的区别、某门课大纲/课程内容、排课时间/日历安排、我该先上哪个）
   - "analytical": 需要结合学员情况做个人规划的查询（如三个月学习计划、转行路径、续费方案、考前冲刺安排）

2. 建议的处理模式:
   - "reactive": 适用于查课、报价、进度、概念解释、课程对比与「先上哪个」等以资料库事实为主的问题；询问某门课「课程内容/大纲」或「排课时间/日历安排」必须走 reactive
   - "deliberative": 仅当明确需要结合已学课程、学习目标与预算做「个人学习计划/路径/课程组合」时使用

请以JSON格式返回结果，包含以下字段:
- query_type: 查询类型（上述三种类型之一）
- processing_mode: 处理模式（上述两种模式之一）
- reasoning: 决策理由的简要说明
"""

# ReAct 决策与最终回答 Prompt（模块级，便于维护）
REACTIVE_AGENT_PROMPT = """你是教培机构咨询规划助手的反应式模块（ReAct）。根据用户问题决定是否调用工具，再回答。

学员信息:
{student_profile}

可用工具（均以资料库文件为准，禁止使用训练记忆中的旧班次数据）：
1. 课程班次查询 — 从资料库汇总班次：价格、上课星期/时间、讲师、报名条件、各期开结课，以及「各课程剩余名额」表中的实时名额（每日一更）
2. 课程内容查询 — 根据课程关键词，从资料库读取某门课的大纲、周次主题、核心内容、学习产出
3. 排课表查询 — 根据课程关键词，从资料库读取某门课的上课星期、时间、各期开结课与具体有课日期
4. 报班活动查询 — 从资料库读取限时优惠 / 续报折扣 / 折上折 / 续报送三门课等活动规则与费用参考

调用格式（单独一行，可写多个）：
[CALL_TOOL:工具名]
例如：
[CALL_TOOL:课程内容查询]
[CALL_TOOL:排课表查询]
[CALL_TOOL:报班活动查询]

何时必须调用工具：
- 问「学什么/课程内容/大纲/每周主题/学习产出」→ 课程内容查询
- 问「什么时候上课/排课/日历/哪天有课/第几期」→ 排课表查询
- 只问价格、有哪些班、班次总览 → 课程班次查询
- 问「限时优惠/限时折扣/打折/几折/折上折/续报福利/赠课/送三门/活动价/现价」→ 必须调用报班活动查询；若还问原价/名额可同时调用课程班次查询
- 同时问内容与时间 → 两个工具都调用

何时不要调用工具：
- 纯闲聊、与课程无关、或仅根据已给学员信息就能答的问题

注意：
- 不要编造资料库中没有的课程、日期、价格、名额、折扣
- 本轮若需要工具，正文可简短说明「正在查询资料库」，并务必带上 CALL_TOOL 标记
- 若不需要工具，直接给出简洁准确回答，不要输出 CALL_TOOL

用户问题: {input}
"""

FINAL_ANSWER_PROMPT = """你是教培机构咨询规划助手。请严格依据「工具返回的资料库原文」回答用户问题，不要编造未出现的信息。

用户问题: {input}
学员信息: {student_profile}

资料库检索结果:
{tool_result}

回答要求：
1. 直接回应用户问的点（内容题讲大纲/周次；时间题讲星期、时段、期次与关键日期；活动题讲折扣/折上折/赠课/参考现价）
2. 只引用检索结果中出现的事实；缺失则明确说「资料库中未写明」
3. 结构清晰：可用分点或按周次/期次列表，避免空话
4. 若命中多门课，先列候选再分别简述，或请用户确认课程名
5. 语气专业友好，面向学员或课程顾问
6. 不要承诺就业、考证通过等结果
7. 若检索结果含报班活动表，优先用其中的折扣规则与费用参考表作答，不要说「资料库没有活动信息」

请输出可直接展示给用户的自然语言回答。
"""

ANALYSIS_PROMPT = """你是一个教培机构咨询规划助手的分析引擎。请仅依据「学员画像」与「资料库事实(kb_facts)」做规划分析，禁止编造课程事实。

用户查询: {user_query}

学员信息:
{student_profile}

已装配的课程与学情数据（含资料库原文）:
{course_data}

分析要求:
1. 学习目标与当前学籍的差距（依据画像中的 learning_goal 与 current_courses）
2. 已报/候选课程是否匹配目标（课程名、价格、先修、课时等硬事实必须来自 kb_facts）
3. 预算是否支撑推荐路径（用画像 budget 与 kb_facts 价格核算；算不出则写「资料库未写明」）
4. 个性化课程顺序与学习建议（软建议可推理，但引用的课名/价格/时间必须来自 kb_facts）
5. 完课、难度和节奏风险（不得编造完课率、就业率等库外统计；时间节奏若用户问题有提及则参考问题，勿臆造周学时约束）
6. 预期能达到的阶段结果（学习阶段表述即可，禁止承诺就业/考证通过/收入）
7. 路径必须遵守 learning_path_note：零基础先 C01；C04 前需数据分析；C06 前建议先大模型。禁止声称「资料库没有基础课」，也禁止编造「速成班/预备班」等库外课名

以JSON格式返回分析结果，包含以下字段:
- goal_gap: 目标与当前学籍的差距
- course_match: 现有课程匹配情况
- recommendations: 课程与学习建议列表（每条尽量标注依据：资料库 / 画像）
- risk_analysis: 完课与节奏风险
- expected_outcomes: 预期学习结果
- fact_basis: 本分析实际用到的关键资料库事实摘要（价格/时间/先修等）
"""

RECOMMENDATION_PROMPT = """你是一个教培机构咨询规划助手。请根据分析结果与资料库事实，为学员准备最终的学习规划。

用户查询: {user_query}

学员信息:
{student_profile}

资料库事实与收集数据:
{course_data}

分析结果:
{analysis_results}

写作要求:
1. 提供清晰、可执行的学习规划，语言友好，避免空泛鼓励
2. 规划应包括：总体策略、行动步骤、课程组合、预算安排、时间框架、后续跟进；上课节奏结合资料库排课与用户问题，勿编造画像中不存在的周学时/长短期约束
3. 硬事实（价格、上课时间、期次、名额、先修、大纲主题）只能来自「资料库事实」；缺失则写「资料库未写明」，禁止补造
4. 软建议以画像中的预算、已学课程、学习目标为主硬约束；可用「依据画像」「依据资料库」区分表述
5. 不要承诺就业、考试通过或收入结果
6. 不要引入资料库中不存在的课程
7. 基础课请明确写「C01 Python入门班」（及资料库中的价格/先修），禁止写「资料库未提供基础课」或「请顾问另查零基础课」；进阶按 learning_path_note 阶梯推荐，勿让零基础直接跨到 C04/C06

返回格式应为自然语言文本，适合直接呈现给学员或课程顾问。
"""

# ==================== 资料库课程索引与解析工具 ====================
KB_DIR = Path(__file__).resolve().parent / "资料库"

# 课程元数据：关键词用于命中用户提问中的课程名
COURSE_KB_META: List[Dict[str, Any]] = [
    {
        "course_id": "C01",
        "name": "Python入门班",
        "keywords": ["Python入门班", "Python入门", "零基础Python", "python入门"],
    },
    {
        "course_id": "C02",
        "name": "Python数据分析班",
        "keywords": ["Python数据分析班", "Python数据分析", "数据分析班", "pandas"],
    },
    {
        "course_id": "C04",
        "name": "机器学习项目班",
        "keywords": ["机器学习项目班", "机器学习", "ML项目班", "YOLO"],
    },
    {
        "course_id": "C05",
        "name": "大模型应用开发班",
        "keywords": [
            "大模型应用开发班", "大模型应用开发课", "大模型应用开发",
            "大模型课", "大模型", "LLM", "RAG", "应用开发课",
        ],
    },
    {
        "course_id": "C06",
        "name": "AI智能体实战班",
        "keywords": ["AI智能体实战班", "AI智能体", "智能体实战", "智能体", "LangGraph", "Agent"],
    },
]


def _resolve_kb_file(course_id: str, kind: Literal["content", "schedule"]) -> Optional[Path]:
    """按课程ID与类型在资料库中定位xlsx（兼容「_新版」等后缀）。"""
    if kind == "content":
        pattern = f"{course_id}_*课程内容简介*.xlsx"
    else:
        pattern = f"{course_id}_*排课表*.xlsx"
    matches = sorted(KB_DIR.glob(pattern))
    return matches[0] if matches else None


def build_course_kb_index() -> List[Dict[str, Any]]:
    """构建课程索引：课程名/关键词 → 内容文件与排课文件路径。"""
    index: List[Dict[str, Any]] = []
    for meta in COURSE_KB_META:
        content_path = _resolve_kb_file(meta["course_id"], "content")
        schedule_path = _resolve_kb_file(meta["course_id"], "schedule")
        index.append({
            **meta,
            "content_path": content_path,
            "schedule_path": schedule_path,
        })
    return index


COURSE_KB_INDEX = build_course_kb_index()
COURSE_NAME_LIST = "、".join(item["name"] for item in COURSE_KB_INDEX)

# 各课程剩余名额统一表（每日一更）
SEATS_XLSX = KB_DIR / "各课程剩余名额.xlsx"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_remaining_seats() -> Dict[str, Any]:
    """
    读取资料库「各课程剩余名额.xlsx」。
    返回 update_date / note / by_course / by_key。
    """
    result: Dict[str, Any] = {
        "update_date": "",
        "note": "本表剩余名额需根据前一天新增的报名名额调整，每日一更。",
        "by_course": {},
        "by_key": {},
    }
    if not SEATS_XLSX.exists():
        return result

    wb = openpyxl.load_workbook(SEATS_XLSX, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    # A2: 数据更新日期：YYYY-MM-DD...
    date_line = _cell_text(ws.cell(2, 1).value)
    if "：" in date_line:
        result["update_date"] = date_line.split("：", 1)[1].split("（", 1)[0].strip()
    elif ":" in date_line:
        result["update_date"] = date_line.split(":", 1)[1].split("（", 1)[0].strip()

    for r in range(5, (ws.max_row or 0) + 1):
        cid = _cell_text(ws.cell(r, 1).value)
        if not cid or not cid.startswith("C"):
            note = _cell_text(ws.cell(r, 1).value)
            if note.startswith("说明"):
                result["note"] = note
            continue
        term = _cell_text(ws.cell(r, 3).value)
        start = _cell_text(ws.cell(r, 4).value)
        end = _cell_text(ws.cell(r, 5).value)
        seats_raw = ws.cell(r, 6).value
        try:
            seats = int(seats_raw)
        except (TypeError, ValueError):
            continue
        entry = {"term": term, "seats": seats, "start": start, "end": end}
        result["by_course"].setdefault(cid, []).append(entry)
        result["by_key"][(cid, term)] = seats
    wb.close()
    return result


REMAINING_SEATS = load_remaining_seats()


def format_seats_for_course(course_id: str) -> str:
    """格式化某门课各期剩余名额文本。"""
    rows = REMAINING_SEATS.get("by_course", {}).get(course_id) or []
    if not rows:
        return "  剩余名额: 资料库中未写明"
    update_date = REMAINING_SEATS.get("update_date") or "未知"
    lines = [f"  剩余名额（数据日期 {update_date}，每日一更）:"]
    for row in rows:
        lines.append(f"  - {row['term']}: 剩余 {row['seats']} 名")
    note = REMAINING_SEATS.get("note") or ""
    if note:
        lines.append(f"  名额说明: {note}")
    return "\n".join(lines)


def match_courses_by_query(query: str) -> List[Dict[str, Any]]:
    """
    根据用户问题匹配课程。
    1) 问题中出现课程全名时，返回所有被点名的课程（支持多课对比）
    2) 否则按关键词最长匹配，定位单门课（同分可多门）
    """
    q = query.lower()

    # 多课对比：凡课程全名出现在问题中的全部纳入（避免只留下最长那一门）
    named = [item for item in COURSE_KB_INDEX if item["name"].lower() in q]
    if named:
        return named

    scored: List[Tuple[int, Dict[str, Any]]] = []
    for item in COURSE_KB_INDEX:
        best_len = 0
        for kw in item["keywords"]:
            if kw and kw.lower() in q:
                best_len = max(best_len, len(kw))
        if best_len > 0:
            scored.append((best_len, item))
    if not scored:
        return []
    max_len = max(s for s, _ in scored)
    return [item for length, item in scored if length == max_len]


def _sheet_table_to_text(ws, start_row: int, max_cols: int = 8) -> str:
    """从指定行起读取连续表格为纯文本（跳过全空行）。"""
    lines: List[str] = []
    for r in range(start_row, (ws.max_row or 0) + 1):
        vals = [_cell_text(ws.cell(r, c).value) for c in range(1, max_cols + 1)]
        if not any(vals):
            continue
        lines.append(" | ".join(vals))
    return "\n".join(lines)


def parse_content_xlsx(path: Path) -> str:
    """解析课程内容简介xlsx：元信息 + 大纲表 + 期次一览摘要。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts: List[str] = [f"【文件】{path.name}"]

    if "课程内容简介" in wb.sheetnames:
        ws = wb["课程内容简介"]
        title = _cell_text(ws.cell(1, 1).value)
        meta = _cell_text(ws.cell(2, 1).value)
        note = _cell_text(ws.cell(3, 1).value)
        if title:
            parts.append(f"标题: {title}")
        if meta:
            parts.append(f"班次信息: {meta}")
        if note:
            parts.append(f"说明: {note}")
        parts.append("大纲（周次 | 单元主题 | 课程核心内容 | 学习产出 | 课型）:")
        parts.append(_sheet_table_to_text(ws, start_row=5, max_cols=5))

    if "期次一览" in wb.sheetnames:
        ws2 = wb["期次一览"]
        parts.append("期次一览（期次 | 开课日期 | 结课日期 | 课节数 | 上课星期 | 上课时间 | 颜色）:")
        parts.append(_sheet_table_to_text(ws2, start_row=3, max_cols=7))

    wb.close()
    return "\n".join(parts)


def parse_schedule_xlsx(path: Path) -> str:
    """
    解析排课表xlsx：排课说明 + 全部课节清单（若无则汇总各月有课日期明细）。
    不读取日历格子，控制上下文长度。
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts: List[str] = [f"【文件】{path.name}"]

    if "排课说明" in wb.sheetnames:
        ws = wb["排课说明"]
        for r in range(1, 4):
            text = _cell_text(ws.cell(r, 1).value)
            if text:
                parts.append(text)
        parts.append("期次总览（期次 | 上课星期 | 开课 | 结课 | 课节数 | 颜色标识）:")
        parts.append(_sheet_table_to_text(ws, start_row=5, max_cols=6))

    if "全部课节清单" in wb.sheetnames:
        ws_all = wb["全部课节清单"]
        parts.append("全部课节清单（期次 | 节次 | 上课日期 | 星期 | 时间 | 讲师 | 颜色）:")
        parts.append(_sheet_table_to_text(ws_all, start_row=3, max_cols=7))
    else:
        # 无汇总表时，从各月「有课日期」明细提取
        detail_blocks: List[str] = []
        for name in wb.sheetnames:
            if name in ("排课说明", "全部课节清单"):
                continue
            ws_m = wb[name]
            header_row = None
            for r in range(1, (ws_m.max_row or 0) + 1):
                if _cell_text(ws_m.cell(r, 1).value) == "日期":
                    header_row = r
                    break
            if header_row is None:
                continue
            table = _sheet_table_to_text(ws_m, start_row=header_row, max_cols=7)
            if table:
                detail_blocks.append(f"## {name}\n{table}")
        if detail_blocks:
            parts.append("各月有课日期明细:")
            parts.extend(detail_blocks)

    wb.close()
    return "\n".join(parts)


def _not_found_message(query: str) -> str:
    return (
        f"未在资料库中找到与「{query}」匹配的课程。\n"
        f"可选课程：{COURSE_NAME_LIST}。\n"
        "请换用完整课程名再问。"
    )


def summarize_course_batch(item: Dict[str, Any]) -> str:
    """从课程内容简介xlsx汇总单班次：元信息 + 期次一览（资料库为唯一数据源）。"""
    path = item.get("content_path")
    if path is None or not Path(path).exists():
        return f"{item['name']}（{item['course_id']}）: 资料库中未找到课程内容简介文件。"

    wb = openpyxl.load_workbook(Path(path), data_only=True, read_only=True)
    lines: List[str] = [f"{item['name']}（{item['course_id']}）"]

    if "课程内容简介" in wb.sheetnames:
        ws = wb["课程内容简介"]
        meta = _cell_text(ws.cell(2, 1).value)
        if meta:
            lines.append(f"  班次信息: {meta}")

    if "期次一览" in wb.sheetnames:
        ws2 = wb["期次一览"]
        # 表头在第3行：期次/开课/结课/...
        term_rows: List[str] = []
        for r in range(4, (ws2.max_row or 0) + 1):
            label = _cell_text(ws2.cell(r, 1).value)
            start = _cell_text(ws2.cell(r, 2).value)
            end = _cell_text(ws2.cell(r, 3).value)
            sessions = _cell_text(ws2.cell(r, 4).value)
            weekdays = _cell_text(ws2.cell(r, 5).value)
            time_slot = _cell_text(ws2.cell(r, 6).value)
            if not label:
                continue
            term_rows.append(
                f"  - {label}: {start} 至 {end}，{sessions}节，{weekdays}，{time_slot}"
            )
        if term_rows:
            lines.append("  期次:")
            lines.extend(term_rows)

    wb.close()
    lines.append(format_seats_for_course(item["course_id"]))
    return "\n".join(lines)


def query_course_schedule(query: str = "") -> str:
    """
    课程班次查询：从资料库汇总班次总览（内容简介 + 剩余名额表）。
    若问题中能命中具体课程则只返回该课；否则返回全部课程。
    """
    matched = match_courses_by_query(query) if query.strip() else []
    targets = matched if matched else COURSE_KB_INDEX

    update_date = REMAINING_SEATS.get("update_date") or "未知"
    seats_note = REMAINING_SEATS.get("note") or ""
    blocks = [
        "在售班次（来源：资料库课程内容简介 + 各课程剩余名额表；名额以名额表为准）:",
        f"名额表更新日期: {update_date}",
    ]
    if seats_note:
        blocks.append(f"名额更新规则: {seats_note}")
    for i, item in enumerate(targets, 1):
        blocks.append(f"{i}. {summarize_course_batch(item)}")

    result = "\n".join(blocks)
    print("[DEBUG] 课程班次查询结果长度:", len(result))
    return result


def query_course_content(query: str = "") -> str:
    """根据关键词从资料库读取课程内容简介。"""
    matched = match_courses_by_query(query)
    if not matched:
        return _not_found_message(query)

    blocks: List[str] = []
    if len(matched) > 1:
        names = "、".join(m["name"] for m in matched)
        blocks.append(f"命中多门课程候选：{names}。以下分别给出资料库内容。")

    for item in matched:
        path = item.get("content_path")
        if path is None or not Path(path).exists():
            blocks.append(f"【{item['name']}】资料库中未找到课程内容简介文件。")
            continue
        blocks.append(f"===== {item['name']}（{item['course_id']}）课程内容 =====")
        blocks.append(parse_content_xlsx(Path(path)))

    result = "\n".join(blocks)
    print("[DEBUG] 课程内容查询结果长度:", len(result))
    return result


def query_course_timetable(query: str = "") -> str:
    """根据关键词从资料库读取排课表（说明+课节明细）。"""
    matched = match_courses_by_query(query)
    if not matched:
        return _not_found_message(query)

    blocks: List[str] = []
    if len(matched) > 1:
        names = "、".join(m["name"] for m in matched)
        blocks.append(f"命中多门课程候选：{names}。以下分别给出排课信息。")

    for item in matched:
        path = item.get("schedule_path")
        if path is None or not Path(path).exists():
            blocks.append(f"【{item['name']}】资料库中未找到排课表文件。")
            continue
        blocks.append(f"===== {item['name']}（{item['course_id']}）排课表 =====")
        blocks.append(parse_schedule_xlsx(Path(path)))

    result = "\n".join(blocks)
    print("[DEBUG] 排课表查询结果长度:", len(result))
    return result


# 报班活动说明表（限时优惠 / 续报福利）；优先用「限时优惠」完整版文件名
PROMO_XLSX_CANDIDATES = [
    KB_DIR / "报班活动-限时优惠与续报福利说明.xlsx",
    KB_DIR / "报班活动-续报福利说明.xlsx",
]


def _resolve_promo_xlsx() -> Optional[Path]:
    """定位资料库中的报班活动说明 xlsx。"""
    for path in PROMO_XLSX_CANDIDATES:
        if path.exists():
            return path
    # 兼容后续改名：匹配「报班活动」前缀
    matches = sorted(KB_DIR.glob("报班活动*.xlsx"))
    return matches[0] if matches else None


def _sheet_to_text_blocks(ws, max_rows: int = 40, max_cols: int = 8) -> List[str]:
    """把工作表前若干行转成可读文本行。"""
    lines: List[str] = []
    for r in range(1, min((ws.max_row or 0), max_rows) + 1):
        cells: List[str] = []
        empty = True
        for c in range(1, min((ws.max_column or 0), max_cols) + 1):
            val = ws.cell(r, c).value
            if val is None or str(val).strip() == "":
                cells.append("")
            else:
                empty = False
                cells.append(str(val).strip())
        if empty:
            continue
        # 去掉尾部空列
        while cells and cells[-1] == "":
            cells.pop()
        lines.append(" | ".join(cells))
    return lines


def query_promo_activity(query: str = "") -> str:
    """
    报班活动查询：读取资料库「报班活动-*.xlsx」中的限时优惠、续报折扣、赠课与费用参考。
    query 仅用于日志；活动规则按整表返回，保证折扣/折上折/赠课完整可答。
    """
    path = _resolve_promo_xlsx()
    if path is None:
        return "资料库中未找到报班活动说明表（报班活动-*.xlsx）。"

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    blocks: List[str] = [
        f"报班活动说明（来源：资料库/{path.name}）",
        "说明：限时优惠活动管学费折扣；续报福利管赠课；二者可同时享受。",
    ]
    # 优先输出与折扣最相关的工作表，其余随后
    preferred = ["活动概览", "限时优惠规则", "续报赠课三门", "费用参考", "课程进阶路径", "顾问话术"]
    sheet_names = list(wb.sheetnames)
    ordered = [n for n in preferred if n in sheet_names]
    ordered += [n for n in sheet_names if n not in ordered]

    # 若问题明显只问折扣/赠课，仍返回全表关键页，避免漏折上折
    for name in ordered:
        ws = wb[name]
        lines = _sheet_to_text_blocks(ws)
        if not lines:
            continue
        blocks.append(f"===== {name} =====")
        blocks.extend(lines)

    wb.close()
    result = "\n".join(blocks)
    print(f"[DEBUG] 报班活动查询: file={path.name}, query={query[:40]!r}, len={len(result)}")
    return result


def extract_tool_calls(text: str) -> List[str]:
    """从LLM输出中解析 [CALL_TOOL:工具名]，按出现顺序去重。"""
    names = re.findall(r"\[CALL_TOOL:([^\]]+)\]", text)
    seen = set()
    ordered: List[str] = []
    for name in names:
        name = name.strip()
        if name and name not in seen:
            seen.add(name)
            ordered.append(name)
    return ordered


def _clip_kb_text(text: str, max_chars: int) -> str:
    """控制写入规划上下文的资料库原文长度。"""
    if not text:
        return ""
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n...(已截断)"


# 规划用路径说明（与「课程路径-基础与进阶.html」一致，写入 course_data 供分析/推荐引用）
LEARNING_PATH_NOTE = (
    "资料库课程阶梯（禁止编造库外课名如「速成班/预备班」）："
    "基础层唯一明确基础课=C01 Python入门班（零基础可报）；"
    "进阶衔接=C02 Python数据分析班（建议具备入门基础）、C05 大模型应用开发班（要求具备入门基础）；"
    "进阶实战=C04 机器学习项目班（要求学过数据分析）、C06 AI智能体实战班（建议先完成大模型应用开发班）。"
    "零基础学员必须先规划 C01，不得声称「资料库未提供基础课」。"
)


def assemble_deliberative_course_data(
    user_query: str,
    student_profile: Optional[Dict[str, Any]],
) -> Dict[str, Any]:
    """
    深思熟虑用数据装配：调用资料库工具写入 kb_facts，不生成模拟学情。
    规划必须拉全量在售班次与内容（避免只命中进阶课导致「找不到基础课」幻觉）。
    排课表：仅当问题涉及日历/冲突/具体有课日期时查询。
    """
    profile = student_profile or {}
    matched = match_courses_by_query(user_query)
    tools_used: List[str] = []
    unknown: List[str] = []

    # 1) 班次与名额：空查询强制全量，保证 C01 等基础课进入上下文
    schedule_text = query_course_schedule("")
    tools_used.append("课程班次查询")
    if not (schedule_text or "").strip():
        unknown.append("课程班次与名额：资料库未返回有效文本")

    # 2) 课程内容：全量拉取，支撑阶梯路径与先修判断
    content_blocks = [query_course_content(item["name"]) for item in COURSE_KB_INDEX]
    content_text = "\n\n".join(content_blocks)
    tools_used.append("课程内容查询")
    if not (content_text or "").strip():
        unknown.append("课程内容：资料库未返回有效文本")

    # 3) 排课明细（按需；未点名课程时同样全量）
    timetable_keys = ("哪天", "日历", "冲突", "排课明细", "有课日期", "上课日期")
    need_timetable = any(k in user_query for k in timetable_keys)
    timetable_text = ""
    if need_timetable:
        if matched:
            timetable_text = query_course_timetable(user_query)
        else:
            timetable_blocks = [query_course_timetable(item["name"]) for item in COURSE_KB_INDEX]
            timetable_text = "\n\n".join(timetable_blocks)
        tools_used.append("排课表查询")
        if not (timetable_text or "").strip():
            unknown.append("排课表：资料库未返回有效文本")
    else:
        timetable_text = "本次未查询排课表（问题未涉及具体日历或上课日期冲突）"

    # 4) 报班活动：涉及折扣/续报赠课/折上折时装配，供规划引用真实活动规则
    promo_keys = (
        "限时", "优惠", "折扣", "打折", "几折", "折上折",
        "续报", "赠课", "送三门", "活动价", "现价", "福利",
    )
    need_promo = any(k in user_query for k in promo_keys)
    promo_text = ""
    if need_promo:
        promo_text = query_promo_activity(user_query)
        tools_used.append("报班活动查询")
        if not (promo_text or "").strip():
            unknown.append("报班活动：资料库未返回有效文本")
    else:
        promo_text = "本次未查询报班活动表（问题未涉及限时优惠/续报折扣/赠课）"

    matched_names = [m["name"] for m in matched] if matched else []
    all_names = [m["name"] for m in COURSE_KB_INDEX]

    return {
        "student_profile": profile,
        "matched_courses": matched_names,
        "catalog_courses": all_names,
        "learning_path_note": LEARNING_PATH_NOTE,
        "tools_used": tools_used,
        "kb_facts": {
            "班次与名额": _clip_kb_text(schedule_text, 8000),
            "课程内容": _clip_kb_text(content_text, 12000),
            "排课表": _clip_kb_text(timetable_text, 8000) if need_timetable else timetable_text,
            "报班活动": _clip_kb_text(promo_text, 8000) if need_promo else promo_text,
        },
        "unknown": unknown,
        "source_note": "课程硬事实仅来自资料库工具原文，非模型编造；规划已加载全部在售课程",
    }


# 第一阶段：情境评估 - 确定查询类型和处理模式
def assess_query(state: EducationAdvisorState) -> EducationAdvisorState:
    print("[DEBUG] 进入节点: assess_query")
    """评估用户查询，确定类型和处理模式"""

    try:
        # 准备提示
        prompt = ChatPromptTemplate.from_template(ASSESSMENT_PROMPT)

        # 构建输入
        input_data = {
            "user_query": state["user_query"],
        }

        # 调用LLM
        chain = prompt | llm | JsonOutputParser()
        result = chain.invoke(input_data)
        print("[DEBUG] LLM评估输出:", result)
        print(f"[DEBUG] 分支判断: processing_mode={result.get('processing_mode', '未知')}, query_type={result.get('query_type', '未知')}")
        # 获取处理模式，确保有值
        processing_mode = result.get("processing_mode", "reactive")
        if processing_mode not in ["reactive", "deliberative"]:
            processing_mode = "reactive"  # 默认使用反应式处理
        # 获取查询类型，确保有值
        query_type = result.get("query_type", "emergency")
        if query_type not in ["emergency", "informational", "analytical"]:
            query_type = "emergency"  # 默认为即时查询
        # ==========================
        # 更新状态
        updated_state = {
            **state,
            "query_type": query_type,
            "processing_mode": processing_mode,
        }
        return updated_state
    except Exception as e:
        return {
            **state,
            "error": f"评估阶段出错: {str(e)}",
            "final_response": "评估查询时发生错误，无法处理您的请求。"
        }

# 反应式处理 - 快速响应课程咨询（支持资料库多工具串行）
def reactive_processing(state: EducationAdvisorState) -> EducationAdvisorState:
    print("[DEBUG] 进入节点: reactive_processing")
    """反应式处理模式：ReAct决策 → 可选串行调用工具 → 结合上下文作答"""
    try:
        # 定义可用工具（资料库工具接收用户原问做关键词匹配）
        tools_map = {
            "课程班次查询": query_course_schedule,
            "课程内容查询": query_course_content,
            "排课表查询": query_course_timetable,
            "报班活动查询": query_promo_activity,
        }

        user_query = state["user_query"]
        student_profile = json.dumps(state.get("student_profile", {}), ensure_ascii=False, indent=2)

        # 第一轮：让LLM决定是否需要工具
        prompt = ChatPromptTemplate.from_template(REACTIVE_AGENT_PROMPT)
        chain = prompt | llm | StrOutputParser()
        first_response = chain.invoke({
            "input": user_query,
            "student_profile": student_profile,
        })
        print("[DEBUG] ReAct首轮输出:", first_response)

        tool_names = extract_tool_calls(first_response)
        # 折扣/续报活动类问题：强制走报班活动表；问现价/多少钱时再带上班次原价
        promo_keys = (
            "限时", "优惠", "折扣", "打折", "几折", "折上折",
            "续报", "赠课", "送三门", "活动价", "现价", "福利",
        )
        price_keys = ("多少钱", "价格", "原价", "现价", "学费", "费用")
        need_promo = any(k in user_query for k in promo_keys)
        if need_promo:
            forced: List[str] = []
            if "报班活动查询" not in tool_names:
                forced.append("报班活动查询")
            if any(k in user_query for k in price_keys) and "课程班次查询" not in tool_names:
                forced.append("课程班次查询")
            if forced:
                tool_names = forced + [n for n in tool_names if n not in forced]
                print(f"[DEBUG] 已强制补充工具: {forced}")

        if tool_names:
            tool_blocks: List[str] = []
            for tool_name in tool_names:
                if tool_name not in tools_map:
                    tool_blocks.append(f"【{tool_name}】未知工具，已跳过。")
                    continue
                print(f"[DEBUG] 调用工具: {tool_name}")
                # 工具均传入用户原问：班次查询可按关键词收窄，未命中则返回全部
                tool_out = tools_map[tool_name](user_query)
                tool_blocks.append(f"### 工具[{tool_name}]返回\n{tool_out}")

            tool_result = "\n\n".join(tool_blocks)
            # 活动类问题：在最终作答前再强调必须引用活动表，避免模型仍说「未写明」
            if need_promo:
                tool_result = (
                    tool_result
                    + "\n\n【作答硬约束】用户在问限时优惠/续报折扣/赠课。"
                    "上方「报班活动查询」返回中已有折扣与费用参考，"
                    "必须据此明确回答几折与参考现价；禁止再说资料库未写明限时优惠。"
                )
            final_prompt = ChatPromptTemplate.from_template(FINAL_ANSWER_PROMPT)
            final_chain = final_prompt | llm | StrOutputParser()
            final_response = final_chain.invoke({
                "input": user_query,
                "student_profile": student_profile,
                "tool_result": tool_result,
            })
        else:
            # 清理可能残留的标记后直接回答
            final_response = re.sub(r"\[CALL_TOOL:[^\]]+\]", "", first_response).strip()

        return {
            **state,
            "final_response": final_response
        }
    except Exception as e:
        return {
            **state,
            "error": f"反应式处理出错: {str(e)}",
            "final_response": "处理您的查询时发生错误，无法提供响应。"
        }

# 数据收集 - 查库装配规划数据（不再由模型模拟 collected_data）
def collect_data(state: EducationAdvisorState) -> EducationAdvisorState:
    print("[DEBUG] 进入节点: collect_data")
    """调用资料库工具装配 course_data，供后续分析与推荐使用"""

    try:
        course_data = assemble_deliberative_course_data(
            user_query=state["user_query"],
            student_profile=state.get("student_profile"),
        )
        print(
            "[DEBUG] 深思熟虑已装配资料库事实:",
            f"tools={course_data.get('tools_used')},",
            f"matched={course_data.get('matched_courses')},",
            f"unknown={course_data.get('unknown')}",
        )

        # 更新状态（保留原有返回结构：course_data + current_phase）
        return {
            **state,
            "course_data": course_data,
            "current_phase": "analyze"
        }
    except Exception as e:
        return {
            **state,
            "error": f"数据收集阶段出错: {str(e)}",
            "current_phase": "collect_data"  # 保持在当前阶段
        }

# 深度分析 - 分析学员目标、时间和已报课程
def analyze_data(state: EducationAdvisorState) -> EducationAdvisorState:
    print("[DEBUG] 进入节点: analyze_data")
    """进行深度学习规划分析"""

    try:
        # 确保必要数据已收集
        if not state.get("course_data"):
            return {
                **state,
                "error": "分析阶段缺少课程与学情数据",
                "current_phase": "collect_data"  # 回到数据收集阶段
            }

        # 准备提示
        prompt = ChatPromptTemplate.from_template(ANALYSIS_PROMPT)

        # 构建输入
        input_data = {
            "user_query": state["user_query"],
            "student_profile": json.dumps(state.get("student_profile", {}), ensure_ascii=False, indent=2),
            "course_data": json.dumps(state.get("course_data", {}), ensure_ascii=False, indent=2)
        }

        # 调用LLM
        chain = prompt | llm | JsonOutputParser()
        result = chain.invoke(input_data)

        # 更新状态
        return {
            **state,
            "analysis_results": result,
            "current_phase": "recommend"
        }
    except Exception as e:
        return {
            **state,
            "error": f"分析阶段出错: {str(e)}",
            "current_phase": "analyze"  # 保持在当前阶段
        }

# 生成建议 - 根据分析结果提供学习规划
def generate_recommendations(state: EducationAdvisorState) -> EducationAdvisorState:
    print("[DEBUG] 进入节点: generate_recommendations")
    """生成学习规划和行动步骤"""

    try:
        # 确保分析结果已存在
        if not state.get("analysis_results"):
            return {
                **state,
                "error": "建议生成阶段缺少分析结果",
                "current_phase": "analyze"  # 回到分析阶段
            }

        # 准备提示（同时传入资料库事实，避免推荐阶段丢失硬事实）
        prompt = ChatPromptTemplate.from_template(RECOMMENDATION_PROMPT)

        # 构建输入
        input_data = {
            "user_query": state["user_query"],
            "student_profile": json.dumps(state.get("student_profile", {}), ensure_ascii=False, indent=2),
            "course_data": json.dumps(state.get("course_data", {}), ensure_ascii=False, indent=2),
            "analysis_results": json.dumps(state.get("analysis_results", {}), ensure_ascii=False, indent=2)
        }

        # 调用LLM
        chain = prompt | llm | StrOutputParser()
        result = chain.invoke(input_data)

        # 更新状态
        return {
            **state,
            "final_response": result,
            "current_phase": "respond"
        }
    except Exception as e:
        return {
            **state,
            "error": f"建议生成阶段出错: {str(e)}",
            "current_phase": "recommend"  # 保持在当前阶段
        }

# 创建智能体工作流
def create_education_advisor_workflow() -> StateGraph:
    """创建教培咨询规划混合智能体工作流"""

    # 创建状态图
    workflow = StateGraph(EducationAdvisorState)

    # 添加节点，每个节点都确保返回完整的状态
    workflow.add_node("assess", assess_query)
    workflow.add_node("reactive", reactive_processing)
    workflow.add_node("collect_data", collect_data)
    workflow.add_node("analyze", analyze_data)
    workflow.add_node("recommend", generate_recommendations)

    # 定义一个显式的响应节点函数
    def respond_function(state: EducationAdvisorState) -> EducationAdvisorState:
        """最终响应生成节点，原样返回状态"""
        # 确保final_response字段有值
        if not state.get("final_response"):
            state = {
                **state,
                "final_response": "无法生成响应。请检查处理流程。",
                "error": state.get("error", "未知错误")
            }
        return state

    workflow.add_node("respond", respond_function)

    # 设置入口点
    workflow.set_entry_point("assess")

    # 添加分支路由
    workflow.add_conditional_edges(
        "assess",
        lambda x: "reactive" if x.get("processing_mode") == "reactive" else "collect_data",
        {
            "reactive": "reactive",
            "collect_data": "collect_data"
        }
    )

    # 添加固定路径边
    workflow.add_edge("reactive", "respond")
    workflow.add_edge("collect_data", "analyze")
    workflow.add_edge("analyze", "recommend")
    workflow.add_edge("recommend", "respond")
    workflow.add_edge("respond", END)

    # 编译工作流
    return workflow.compile()

# 示例学员画像：以预算、已学课程、学习目标为主硬约束
SAMPLE_STUDENT_PROFILES = {
    "student1": {
        "student_id": "S2026001",
        "current_level": "零基础",
        "learning_goal": "转行，从Python入门打基础并可继续进阶",
        "budget": 10000.0,
        "current_courses": {}
    },
    "student2": {
        "student_id": "S2026002",
        "current_level": "入门",
        "learning_goal": "在入门基础上评估数据分析并衔接大模型应用开发",
        "budget": 18000.0,
        "current_courses": {
            "Python入门班": "已完成"
        }
    },
    # 资深：已学完入门+数据分析，大模型在读，预算用于继续深造（如机器学习/智能体）
    "student3": {
        "student_id": "S2026003",
        "current_level": "进阶",
        "learning_goal": "在已有基础上继续深造，规划机器学习与AI智能体学习节奏",
        "budget": 10000.0,
        "current_courses": {
            "Python入门班": "已完成",
            "Python数据分析班": "已完成",
            "大模型应用开发班": "进行中",
        }
    },
}

# 交互菜单展示用标签（固定示例；选 0 可自定义任意学籍组合）
STUDENT_MENU_OPTIONS = [
    ("student1", "零基础学员（尚未学习任何课程，目标转行，预算约10000元）"),
    ("student2", "入门学员（Python入门班已完成，预算约18000元）"),
    ("student3", "资深学员（入门与数据分析已完成，大模型进行中，预算约10000元）"),
]


def _parse_course_name_list(raw: str) -> List[str]:
    """解析逗号/顿号分隔的课程名列表。"""
    text = (raw or "").strip()
    if not text:
        return []
    parts = re.split(r"[,，、;；]+", text)
    return [p.strip() for p in parts if p.strip()]


def prompt_custom_student_profile() -> Dict[str, Any]:
    """
    交互录入自定义学员画像，用于覆盖「固定示例」未涵盖的学籍组合。
    课程名建议与资料库一致，例如：Python入门班、Python数据分析班、大模型应用开发班等。
    """
    print("\n--- 自定义学员画像 ---")
    print(f"可选课程参考: {COURSE_NAME_LIST}")

    goal = input("学习目标（如：继续深造机器学习）: ").strip()
    budget_raw = input("预算（元，默认10000）: ").strip()
    budget = float(budget_raw) if budget_raw else 10000.0

    completed = _parse_course_name_list(
        input("已完成课程（多个用逗号分隔，无可回车跳过）: ")
    )
    in_progress = _parse_course_name_list(
        input("进行中课程（多个用逗号分隔，无可回车跳过）: ")
    )

    current_courses: Dict[str, str] = {}
    for name in completed:
        current_courses[name] = "已完成"
    for name in in_progress:
        current_courses[name] = "进行中"

    if not current_courses:
        level = "零基础"
    elif any(v == "进行中" for v in current_courses.values()) or len(current_courses) >= 2:
        level = "进阶"
    else:
        level = "入门"

    return {
        "student_id": "CUSTOM",
        "current_level": level,
        "learning_goal": goal or "继续深造",
        "budget": budget,
        "current_courses": current_courses,
    }


# 运行智能体
def run_education_advisor(
    user_query: str,
    student_id: str = "student1",
    student_profile: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """运行教培咨询规划智能体并返回结果；可传入自定义 student_profile 覆盖示例画像。"""

    # 创建工作流
    agent = create_education_advisor_workflow()

    # 获取学员画像：优先使用调用方传入的自定义画像
    if student_profile is None:
        student_profile = SAMPLE_STUDENT_PROFILES.get(
            student_id, SAMPLE_STUDENT_PROFILES["student1"]
        )
    else:
        student_id = student_profile.get("student_id", student_id) or "custom"

    # 准备初始状态
    initial_state = {
        "user_query": user_query,
        "student_profile": student_profile,
        "query_type": None,
        "processing_mode": None,
        "course_data": None,
        "analysis_results": None,
        "final_response": None,
        "current_phase": "assess",
        "error": None
    }

    try:
        print("LangGraph Mermaid流程图：")
        print(agent.get_graph().draw_mermaid())

        # 准备 Langfuse 配置（如果启用）
        config = {}
        if LANGFUSE_ENABLED:
            # 通过 CallbackHandler 将 LangGraph/LangChain 执行链路上报到 Langfuse
            langfuse_handler = CallbackHandler()
            config = {
                "callbacks": [langfuse_handler],
                "tags": [
                    "education-advisor",
                    "hybrid-agent",
                    f"student-{student_id}",
                    student_profile.get("current_level", "unknown")
                ],
                "metadata": {
                    "student_id": student_id,
                    "current_level": student_profile.get("current_level", "unknown"),
                    "learning_goal": student_profile.get("learning_goal", "unknown"),
                    "budget": student_profile.get("budget", 0),
                    "user_query": user_query[:100],  # 只记录前100个字符
                    "timestamp": datetime.now().isoformat()
                },
                "run_name": f"education-advisor-{student_id}-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
            }

        # 运行智能体并捕获可能的异常
        if config:
            result = agent.invoke(initial_state, config=config)
        else:
            result = agent.invoke(initial_state)

        # 短生命周期脚本需主动 flush，确保追踪数据发送完成
        if LANGFUSE_ENABLED and langfuse_client is not None:
            langfuse_client.flush()

        return result
    except Exception as e:
        error_msg = str(e)
        print(f"捕获异常: {error_msg}")
        if LANGFUSE_ENABLED and langfuse_client is not None:
            langfuse_client.flush()
        # 返回带有错误信息的状态
        return {
            **initial_state,
            "error": f"执行过程中发生错误: {error_msg}",
            "final_response": "很抱歉，处理您的请求时出现了问题。"
        }

# 主函数
if __name__ == "__main__":
    print("=== 混合智能体 - 教培机构咨询规划助手 ===\n")
    print("使用模型：Qwen-Turbo-Latest\n")

    # 显示 Langfuse 状态
    if LANGFUSE_ENABLED:
        print("Langfuse 追踪已启用")
        print(f"  地址: {LANGFUSE_BASE_URL}")
        print(f"  查看追踪: {LANGFUSE_BASE_URL}\n")
    else:
        print("Langfuse 追踪未启用")
        print("  要启用追踪，请设置环境变量:")
        print("  - LANGFUSE_PUBLIC_KEY: 项目公钥 pk-lf-...")
        print("  - LANGFUSE_SECRET_KEY: 项目私钥 sk-lf-...")
        print("  - LANGFUSE_BASE_URL: http://localhost:3000（本地默认）\n")

    print("-"*50 + "\n")

    # 示例查询：规划类不写死学籍/预算，避免与所选学员画像冲突（约束以画像为准）
    SAMPLE_QUERIES = [
        # 即时/简单查询 - 适合反应式处理
        "Python数据分析班多少钱，每周什么时候上课，还剩几个名额？",
        "大模型应用开发班和AI智能体实战班的价格、上课时间和剩余名额分别是多少？",
        "Python入门班、Python数据分析班和大模型应用开发班有什么区别，我该先上哪个？",

        # 规划性查询 - 适合深思熟虑处理（学籍/预算看所选学员）
        "请结合我的学员信息，帮我规划接下来的学习路径和课程组合。",
        "请结合我的预算和已学课程，评估下一步该怎么选课、如何衔接。",
        "请结合我的情况，评估机器学习项目班和AI智能体实战班，并给出学习节奏建议。",
    ]

    # 用户选择查询示例或输入自定义查询
    print("请选择一个示例查询或输入您自己的查询:")
    print("（规划类问题4-6请以所选学员画像为准，勿在问题里再写「零基础/已学完某课」以免冲突）\n")
    for i, query in enumerate(SAMPLE_QUERIES, 1):
        print(f"{i}. {query}")
    print("0. 输入自定义查询")

    choice = input("\n请输入选项数字(0-6): ")

    if choice == "0":
        user_query = input("请输入您的查询: ")
    else:
        try:
            idx = int(choice) - 1
            if 0 <= idx < len(SAMPLE_QUERIES):
                user_query = SAMPLE_QUERIES[idx]
            else:
                print("无效选择，使用默认查询")
                user_query = SAMPLE_QUERIES[0]
        except ValueError:
            print("无效输入，使用默认查询")
            user_query = SAMPLE_QUERIES[0]

    # 选择学员（固定示例 1-3，或 0 自定义任意学籍）
    print("\n请选择学员:")
    for i, (_, label) in enumerate(STUDENT_MENU_OPTIONS, 1):
        print(f"  {i}. {label}")
    print("  0. 自定义学员（自行填写目标/预算/已完成与进行中课程）")
    student_choice = input(
        f"选择学员 (0-{len(STUDENT_MENU_OPTIONS)}，默认1): "
    ).strip()

    student_id = "student1"
    custom_profile: Optional[Dict[str, Any]] = None
    menu_label = STUDENT_MENU_OPTIONS[0][1]
    if student_choice == "0":
        custom_profile = prompt_custom_student_profile()
        student_id = "custom"
        menu_label = "自定义学员"
    else:
        try:
            choice_num = int(student_choice) if student_choice else 1
            if 1 <= choice_num <= len(STUDENT_MENU_OPTIONS):
                student_id = STUDENT_MENU_OPTIONS[choice_num - 1][0]
                menu_label = STUDENT_MENU_OPTIONS[choice_num - 1][1]
            else:
                print("无效选择，使用学员1")
        except ValueError:
            print("无效输入，使用学员1")

    profile = custom_profile if custom_profile is not None else SAMPLE_STUDENT_PROFILES[student_id]
    courses = profile.get("current_courses") or {}
    course_desc = "尚未学习任何课程" if not courses else "；".join(f"{k}（{v}）" for k, v in courses.items())
    print(f"\n用户查询: {user_query}")
    print(f"选择学员: {menu_label}")
    print(f"学员概况: {profile.get('current_level', '')} / 目标{profile.get('learning_goal', '')} / 预算{profile['budget']}元")
    print(f"在学情况: {course_desc}")
    print("\n正在处理...\n")

    try:
        # 运行智能体
        start_time = datetime.now()
        result = run_education_advisor(
            user_query,
            student_id,
            student_profile=custom_profile,
        )
        end_time = datetime.now()

        # 如果有错误，显示错误信息并退出
        if result.get("error"):
            print(f"处理过程中发生错误: {result['error']}")
            print(f"\n最终响应: {result.get('final_response', '未能生成响应')}")
            process_time = (end_time - start_time).total_seconds()
            print(f"\n处理用时: {process_time:.2f}秒")
            exit(1)

        # 显示处理模式
        process_mode = result.get("processing_mode", "未知")
        if process_mode == "reactive":
            print("【处理模式: 反应式】- 快速响应课程咨询")
        else:
            print("【处理模式: 深思熟虑】- 深度学习规划")

        # 显示结果
        print("\n=== 响应结果 ===\n")
        print(result.get("final_response", "未生成响应"))

        # 显示处理时间
        process_time = (end_time - start_time).total_seconds()
        print(f"\n处理用时: {process_time:.2f}秒")

    except Exception as e:
        print(f"\n运行过程中发生意外错误: {str(e)}")
