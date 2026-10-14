# -*- coding: utf-8 -*-
"""
拆分生成：各班次课程内容简介 + 日历式排课表 + 统一剩余名额表。
规则：
- 期次命名为「2026年第X期」（可滚动多期，不限两期）
- 一周两节：周二/周四/周五中选2天，或周三+周五
- 日历格子：有课只标颜色，不写额外文字
- 在2026-09至2027-01窗口内连续滚动开班，避免月份空档
- 「各课程剩余名额.xlsx」统一存放各期剩余名额，需按前一天新增报名调整，每日一更
"""

from __future__ import annotations

import calendar
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent
calendar.setfirstweekday(calendar.MONDAY)

WINDOW_START = date(2026, 9, 1)
WINDOW_END = date(2027, 1, 31)

TITLE_FONT = Font(bold=True, name="微软雅黑", size=14, color="1F4E79")
SUB_FONT = Font(name="微软雅黑", size=10, color="666666")
HEADER_FONT = Font(bold=True, name="微软雅黑", size=11, color="FFFFFF")
NORMAL_FONT = Font(name="微软雅黑", size=10)
DATE_FONT = Font(name="微软雅黑", size=14, bold=True)
DATE_MUTED = Font(name="微软雅黑", size=12, color="999999")
SMALL_FONT = Font(name="微软雅黑", size=9)
WEEKDAY_FONT = Font(bold=True, name="微软雅黑", size=10)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
GRAY_FILL = PatternFill("solid", fgColor="F2F2F2")
HOLIDAY_FILL = PatternFill("solid", fgColor="FCE4D6")
ALT_FILL = PatternFill("solid", fgColor="F5F9FC")

# 多期颜色轮换：黄/绿/蓝/橙/紫
TERM_COLORS = [
    ("黄", PatternFill("solid", fgColor="FFE699")),
    ("绿", PatternFill("solid", fgColor="C6EFCE")),
    ("蓝", PatternFill("solid", fgColor="BDD7EE")),
    ("橙", PatternFill("solid", fgColor="F8CBAD")),
    ("紫", PatternFill("solid", fgColor="D9D2E9")),
]

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)

WEEKDAY_CN = ["一", "二", "三", "四", "五", "六", "日"]
NATIONAL_DAY = (date(2026, 10, 1), date(2026, 10, 7))

COURSES = [
    {
        "id": "C01",
        "name": "Python入门班",
        "file_stub": "C01_Python入门班",
        "price": 1280,
        "weeks": 4,
        "weekdays": (1, 3),  # 周二、周四
        "time": "19:30-21:30",
        "teacher": "陈老师",
        "condition": "零基础可报",
        "first_start": date(2026, 9, 1),
    },
    {
        "id": "C02",
        "name": "Python数据分析班",
        "file_stub": "C02_Python数据分析班",
        "price": 3680,
        "weeks": 8,
        "weekdays": (2, 4),  # 周三、周五
        "time": "19:30-21:30",
        "teacher": "周老师",
        "condition": "建议具备Python入门基础",
        "first_start": date(2026, 9, 2),
    },
    {
        "id": "C04",
        "name": "机器学习项目班",
        "file_stub": "C04_机器学习项目班",
        "price": 6980,
        "weeks": 10,
        "weekdays": (3, 4),  # 周四、周五
        "time": "19:30-21:30",
        "teacher": "赵老师",
        "condition": "要求已完成Python数据分析",
        "first_start": date(2026, 9, 3),
    },
    {
        "id": "C05",
        "name": "大模型应用开发班",
        "file_stub": "C05_大模型应用开发班",
        "price": 5280,
        "weeks": 8,
        "weekdays": (1, 3),  # 周二、周四
        "time": "19:30-21:30",
        "teacher": "吴老师",
        "condition": "要求具备Python入门基础",
        "first_start": date(2026, 9, 1),
    },
    {
        "id": "C06",
        "name": "AI智能体实战班",
        "file_stub": "C06_AI智能体实战班",
        "price": 7580,
        "weeks": 6,
        "weekdays": (2, 4),  # 周三、周五
        "time": "19:30-21:30",
        "teacher": "孙老师",
        "condition": "含工具调用、LangGraph与工作流编排；建议先完成大模型应用开发班",
        "first_start": date(2026, 9, 2),
    },
]

CONTENTS = {
    "C01": [
        ["第1周", "环境与语法基础", "Python安装与虚拟环境、变量与数据类型、条件与循环", "完成本地环境搭建；能写基础分支与循环", "实操"],
        ["第1周", "函数与模块", "函数定义、参数与返回值、常用标准库、代码规范", "能把重复逻辑封装成函数", "实操"],
        ["第2周", "数据结构", "列表、字典、集合、字符串处理、文件读写", "完成文本与CSV读写小练习", "实操"],
        ["第2周", "异常与调试", "异常处理、日志打印、用Cursor辅助排查错误", "能独立定位并修复常见报错", "实操+工具"],
        ["第3周", "面向对象入门", "类与对象、简单封装、项目目录结构", "完成一个小型业务对象建模", "实操"],
        ["第3周", "第三方库与API初识", "pip、requests、JSON解析；为后续大模型API打基础", "能调用公开HTTP接口并解析结果", "实操"],
        ["第4周", "综合实战", "CASE：多张Excel报表汇总；简易命令行工具", "交付可运行小项目与说明文档", "项目"],
        ["第4周", "AI编程加速", "Cursor Rules基础、用AI补全与重构入门代码", "掌握用AI辅助写Python的基本流程", "实操"],
    ],
    "C02": [
        ["第1周", "数据分析环境", "Jupyter、Pandas与Numpy概览", "搭好分析环境并完成首个Notebook", "实操"],
        ["第1周", "数据读取与清洗", "CSV/Excel读取、缺失值、重复值、类型转换", "完成一份原始表清洗", "实操"],
        ["第2周", "索引与筛选", "loc/iloc、条件筛选、分组聚合", "按业务条件产出汇总表", "实操"],
        ["第2周", "多表关联", "merge/join、concat、主键设计", "完成订单-用户多表关联", "实操"],
        ["第3周", "时间序列基础", "日期解析、重采样、同比环比", "做出周/月趋势汇总", "实操"],
        ["第3周", "可视化", "Matplotlib/Seaborn常用图、图表叙事", "输出可讲解的分析图", "实操"],
        ["第4周", "业务指标体系", "转化漏斗、留存、RFM分层思路", "为给定数据集设计指标口径", "理论+实操"],
        ["第4周", "SQL与Python协同", "常用SQL与Pandas对照", "同一问题分别用SQL和Pandas完成", "实操"],
        ["第5周", "特征准备入门", "编码、标准化、训练测试划分", "为机器学习准备特征表", "实操"],
        ["第5周", "统计分析常识", "分布、相关、结果可信度直觉", "能解释分析结果是否靠谱", "理论"],
        ["第6周", "自动化报表", "脚本导出、openpyxl美化", "生成周报Excel模板", "项目"],
        ["第6周", "AI辅助分析", "用大模型解读异常波动", "完成一次AI辅助解读演练", "实操"],
        ["第7周", "综合项目A", "电商销售分析全流程", "提交分析报告", "项目"],
        ["第7周", "项目讲评", "结论表达与常见坑复盘", "根据讲评修订报告", "讲评"],
        ["第8周", "综合项目B", "运营活动复盘看板", "提交可复用分析模板", "项目"],
        ["第8周", "结业与衔接", "通向机器学习/大模型班的路径", "明确后续可选班次", "总结"],
    ],
    "C04": [
        ["第1周", "机器学习全景", "监督/无监督、训练验证测试、过拟合", "说清建模标准流程", "理论"],
        ["第1周", "特征工程实战", "编码、归一化、特征选择、防泄漏", "完成特征流水线", "实操"],
        ["第2周", "经典模型", "线性/逻辑回归、树模型、评估指标", "完成二分类基线", "实操"],
        ["第2周", "调参与对比", "交叉验证、搜索调参、对比表", "输出模型对比报告", "实操"],
        ["第3周", "神经网络基础", "激活、损失、反向传播、梯度下降", "建立最小网络直觉", "理论+实操"],
        ["第3周", "PyTorch入门", "张量、自动求导、DataLoader", "跑通首个训练循环", "实操"],
        ["第4周", "训练技巧", "学习率、正则、早停、保存加载", "稳定复现一次训练", "实操"],
        ["第4周", "非视觉任务", "表格MLP/简单文本分类", "完成一个小项目", "项目"],
        ["第5周", "计算机视觉基础", "分类/检测/分割与标注格式", "读懂检测数据目录", "理论"],
        ["第5周", "YOLO体系", "YOLO演进与ultralytics工具链", "完成官方示例推理", "实操"],
        ["第6周", "自定义检测", "标注→训练→mAP评估", "训练小样本检测模型", "项目"],
        ["第6周", "缺陷分析调优", "难例、增强、阈值", "提交调优记录", "实操"],
        ["第7周", "多模态对比", "YOLO vs Qwen-VL零样本质检思路", "完成对比实验纪要", "项目"],
        ["第7周", "部署初步", "导出与本地推理接口", "封装可调用服务", "实操"],
        ["第8周", "综合项目启动", "选题、里程碑、数据范围", "确定项目计划", "项目"],
        ["第8周", "数据工程冲刺", "EDA、划分、基线", "提交基线指标", "项目"],
        ["第9周", "模型迭代", "误差分析驱动第二轮训练", "记录指标提升原因", "项目"],
        ["第9周", "复盘可视化", "混淆矩阵、失败案例库", "完成复盘文档", "项目"],
        ["第10周", "结业答辩", "演示、指标、局限与迭代", "完成结业答辩", "答辩"],
        ["第10周", "进阶路径", "转向大模型/智能体班的条件", "给出个人路径建议", "总结"],
    ],
    "C05": [
        ["第1周", "大模型原理", "生成式AI、Token、Temperature/TopP、选型", "解释关键推理参数影响", "理论"],
        ["第1周", "API实战", "DashScope/Qwen、系统/用户提示词", "跑通情感分析CASE", "实操"],
        ["第2周", "提示与上下文工程", "高级提示、长上下文、结构化输出", "沉淀可复用提示模板", "理论+实操"],
        ["第2周", "Function Calling入门", "工具调用、JSON约束", "完成单工具调用Demo", "实操"],
        ["第3周", "RAG核心", "三种范式、NativeRAG、私有知识库", "画出RAG架构图", "理论"],
        ["第3周", "向量与检索", "Embedding选型、Faiss入门", "完成文档向量化入库", "实操"],
        ["第4周", "RAG搭建", "本地知识库问答", "交付可演示RAG", "项目"],
        ["第4周", "知识库治理", "Query改写、版本与健康度", "补齐问题清单", "实操"],
        ["第5周", "混合检索", "BM25+向量、多路召回、Rerank", "对比检索效果", "实操"],
        ["第5周", "RAG评估", "检索/生成调试与基础指标", "输出一页评估表", "理论+实操"],
        ["第6周", "LangChain入门", "组件与LCEL任务链", "搭一条可运行链", "实操"],
        ["第6周", "框架内客服", "本地知识客服与ReAct理解", "完成客服Demo", "项目"],
        ["第7周", "多模态数据", "PDF/Word/网页解析与图片理解", "接入一种文档解析", "实操"],
        ["第7周", "综合项目", "FAQ助手：解析→检索→生成→前端", "提交可演示项目", "项目"],
        ["第8周", "效果与成本", "幻觉、引用溯源、成本核算", "给出上线检查清单", "理论"],
        ["第8周", "结业与进阶", "衔接到AI智能体实战班", "明确是否具备Agent班基础", "总结"],
    ],
    "C06": [
        ["第1周", "Agent范式", "反应式/深思熟虑/混合式；CoT与ReAct", "说明三类Agent适用场景", "理论"],
        ["第1周", "工具调用", "Function Calling与MCP差异；自定义工具", "完成多工具注册调用", "实操"],
        ["第2周", "MCP与外部能力", "Host/Client/Server；Fetch/搜索；最小MCP服务", "跑通MCP接入", "实操"],
        ["第2周", "状态与记忆", "对话状态、Memory、追踪注意点", "设计Agent状态字段", "理论+实操"],
        ["第3周", "LangGraph工作流", "节点、边、条件分支；三层架构落地", "实现咨询助手StateGraph", "实操"],
        ["第3周", "人机协同", "审批节点、干预与重试", "加入人工确认节点", "实操"],
        ["第4周", "能力优化", "反馈、业务指标、评估概念", "制定效果评估表", "理论"],
        ["第4周", "可观测性", "本地Langfuse调试与链路打点", "完成一次追踪演示", "实操"],
        ["第5周", "OpenManus解析", "编排/Agent/Memory/Tools生命周期", "本地跑通主流程", "项目"],
        ["第5周", "场景定制", "角色定制与中文咨询流改造", "改造业务向助手", "项目"],
        ["第6周", "多Agent协作", "主Agent调度、拆解分发、结果聚合", "完成双Agent Demo", "项目"],
        ["第6周", "结业答辩", "演示、评估、风险与迭代", "完成结业答辩", "答辩"],
    ],
}


def is_holiday(d: date) -> bool:
    return NATIONAL_DAY[0] <= d <= NATIONAL_DAY[1]


def term_label(d0: date, seq_in_year: dict[int, int]) -> str:
    """按开课年份累计期号：2026年第X期 / 2027年第X期。"""
    year = d0.year
    seq_in_year[year] = seq_in_year.get(year, 0) + 1
    return f"{year}年第{seq_in_year[year]}期"


def weekday_desc(weekdays: tuple[int, ...]) -> str:
    names = "、".join(f"周{WEEKDAY_CN[w]}" for w in weekdays)
    return f"每周{names}（一周两节）"


def color_for_term_index(term_index: int) -> tuple[str, PatternFill]:
    return TERM_COLORS[(term_index - 1) % len(TERM_COLORS)]


def first_class_day(start: date, weekdays: tuple[int, ...]) -> date:
    cur = max(start, WINDOW_START)
    for _ in range(21):
        if cur.weekday() in weekdays and not is_holiday(cur):
            return cur
        cur += timedelta(days=1)
    return cur


def collect_term_dates(start: date, weekdays: tuple[int, ...], weeks: int) -> list[date]:
    need = weeks * 2
    dates: list[date] = []
    cur = first_class_day(start, weekdays)
    while len(dates) < need and cur <= WINDOW_END + timedelta(days=60):
        if cur.weekday() in weekdays and not is_holiday(cur):
            dates.append(cur)
        cur += timedelta(days=1)
    # 只保留窗口内的完整期；若结课超出窗口则整期不要（后面滚动控制）
    return dates


def build_rolling_terms(course: dict) -> list[dict]:
    """
    在窗口内连续滚动开班：上一期上满（周数*2节）后再开下一期，期次不重叠。
    保证「4周课 = 8节课」与课程内容简介一致，不会被后一期挤掉后半段。
    """
    terms = []
    seq_in_year: dict[int, int] = {}
    term_index = 0
    need = course["weeks"] * 2
    cursor = course["first_start"]

    while cursor <= WINDOW_END:
        dates = collect_term_dates(cursor, course["weekdays"], course["weeks"])
        if not dates:
            break
        # 必须凑齐完整课节，且全部落在窗口内
        if len(dates) < need or dates[-1] > WINDOW_END:
            break
        term_index += 1
        label = term_label(dates[0], seq_in_year)
        color_name, fill = color_for_term_index(term_index)
        terms.append(
            {
                "term_index": term_index,
                "label": label,
                "dates": dates,
                "color_name": color_name,
                "fill": fill,
            }
        )
        # 下一期从上一期结课日的次日开始
        cursor = dates[-1] + timedelta(days=1)

    return terms


def style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def write_content_file(course: dict, terms: list[dict]) -> Path:
    path = OUT_DIR / f"{course['file_stub']}_课程内容简介.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "课程内容简介"

    ws["A1"] = f"{course['name']} · 课程内容简介"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:E1")

    ws["A2"] = (
        f"班次编号：{course['id']}　|　价格：{course['price']}元　|　"
        f"{weekday_desc(course['weekdays'])}　|　每节 {course['time']}　|　"
        f"共{course['weeks']}周/{course['weeks']*2}节　|　讲师：{course['teacher']}　|　"
        f"报名：{course['condition']}"
    )
    ws["A2"].font = SUB_FONT
    ws["A2"].alignment = LEFT
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 36

    ws["A3"] = (
        f"排课窗口2026-09至2027-01：一期上满{course['weeks']}周/{course['weeks']*2}节后再开下一期，"
        f"共安排{len(terms)}期。日历有课只标颜色。详见同目录排课表。"
    )
    ws["A3"].font = SUB_FONT
    ws["A3"].alignment = LEFT
    ws.merge_cells("A3:E3")

    headers = ["周次", "单元主题", "课程核心内容", "学习产出", "课型"]
    for i, h in enumerate(headers, 1):
        ws.cell(5, i, h)
    style_header_row(ws, 5, 5)

    for i, row in enumerate(CONTENTS[course["id"]]):
        for j, val in enumerate(row, 1):
            cell = ws.cell(6 + i, j, val)
            cell.font = NORMAL_FONT
            cell.alignment = LEFT
            cell.border = THIN
            if i % 2 == 1:
                cell.fill = ALT_FILL
        ws.row_dimensions[6 + i].height = 42
    for i, w in enumerate([10, 20, 70, 36, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("期次一览")
    ws2["A1"] = f"{course['name']} · 期次一览（滚动多期）"
    ws2["A1"].font = TITLE_FONT
    ws2.merge_cells("A1:G1")
    for i, h in enumerate(["期次", "开课日期", "结课日期", "课节数", "上课星期", "上课时间", "颜色"], 1):
        ws2.cell(3, i, h)
    style_header_row(ws2, 3, 7)
    for idx, term in enumerate(terms):
        values = [
            term["label"],
            term["dates"][0].isoformat(),
            term["dates"][-1].isoformat(),
            len(term["dates"]),
            weekday_desc(course["weekdays"]),
            course["time"],
            term["color_name"],
        ]
        for j, v in enumerate(values, 1):
            cell = ws2.cell(4 + idx, j, v)
            cell.font = NORMAL_FONT
            cell.alignment = CENTER
            cell.border = THIN
            cell.fill = term["fill"]
    for i, w in enumerate([16, 14, 14, 10, 28, 14, 8], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    return safe_save(wb, path)


def build_date_map(terms: list[dict]) -> dict[date, dict]:
    """同一天若多期重叠，保留期号更大的（后开班）。"""
    mapping: dict[date, dict] = {}
    for term in terms:
        for d in term["dates"]:
            mapping[d] = term
    return mapping


def write_calendar_month(ws, start_row: int, year: int, month: int, date_map: dict, course: dict, terms: list[dict]) -> int:
    ws.cell(start_row, 1, f"{year}年{month}月").font = TITLE_FONT
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=7)

    # 图例：本班次实际用到的期次颜色
    ws.cell(start_row, 9, "图例").font = WEEKDAY_FONT
    for i, term in enumerate(terms):
        cell = ws.cell(start_row + (i // 6), 10 + (i % 6), term["label"])
        cell.fill = term["fill"]
        cell.font = SMALL_FONT
        cell.border = THIN
        cell.alignment = CENTER

    header_row = start_row + 1
    for i, name in enumerate(WEEKDAY_CN, 1):
        cell = ws.cell(header_row, i, name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN

    weeks = calendar.monthcalendar(year, month)
    r = header_row + 1
    for week in weeks:
        for col, day in enumerate(week, 1):
            cell = ws.cell(r, col)
            cell.border = THIN
            cell.alignment = CENTER
            if day == 0:
                cell.fill = GRAY_FILL
                cell.value = ""
                continue
            d = date(year, month, day)
            # 只写日期数字，有课只靠底色识别
            cell.value = day
            if d in date_map:
                cell.fill = date_map[d]["fill"]
                cell.font = DATE_FONT
            else:
                cell.font = DATE_MUTED
                if is_holiday(d):
                    cell.fill = HOLIDAY_FILL
        ws.row_dimensions[r].height = 36
        r += 1

    detail_row = r + 1
    ws.cell(detail_row, 1, f"{year}年{month}月有课日期").font = WEEKDAY_FONT
    ws.merge_cells(start_row=detail_row, start_column=1, end_row=detail_row, end_column=7)
    detail_row += 1
    for i, h in enumerate(["日期", "星期", "时间", "期次", "节次序号", "讲师", "备注"], 1):
        cell = ws.cell(detail_row, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN

    # 按完整期次列出本月课节（节次号相对该期全部课表，不被其他期覆盖）
    month_items = []
    for term in terms:
        for sess_no, d in enumerate(term["dates"], 1):
            if d.year == year and d.month == month:
                month_items.append((d, term, sess_no))
    month_items.sort(key=lambda x: (x[0], x[1]["term_index"]))

    detail_row += 1
    if not month_items:
        ws.cell(detail_row, 1, "本月无排课").font = SUB_FONT
        detail_row += 1
    else:
        for d, term, sess_no in month_items:
            values = [
                d.isoformat(),
                f"周{WEEKDAY_CN[d.weekday()]}",
                course["time"],
                term["label"],
                f"第{sess_no}节",
                course["teacher"],
                f"本期共{len(term['dates'])}节/{course['weeks']}周",
            ]
            for c, v in enumerate(values, 1):
                cell = ws.cell(detail_row, c, v)
                cell.font = NORMAL_FONT
                cell.alignment = CENTER
                cell.border = THIN
                cell.fill = term["fill"]
            detail_row += 1
    return detail_row + 2


def write_schedule_file(course: dict, terms: list[dict]) -> Path:
    path = OUT_DIR / f"{course['file_stub']}_排课表.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "排课说明"

    ws["A1"] = f"{course['name']} · 排课表（日历）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")
    ws["A2"] = (
        f"班次 {course['id']}｜{weekday_desc(course['weekdays'])}｜每节 {course['time']}｜"
        f"共{course['weeks']}周/{course['weeks']*2}节｜讲师{course['teacher']}｜"
        f"一期上满{course['weeks']}周/{course['weeks']*2}节后再开下期，共{len(terms)}期。"
        f"日历有课只标颜色，不写额外文字。"
    )
    ws["A2"].font = SUB_FONT
    ws["A2"].alignment = LEFT
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 40

    ws["A3"] = "覆盖：2026年9月—2027年1月。国庆（10/1-10/7）自动避让。期次不重叠、不限两期。"
    ws["A3"].font = SUB_FONT
    ws.merge_cells("A3:G3")

    for i, h in enumerate(["期次", "上课星期", "开课", "结课", "课节数", "颜色标识"], 1):
        ws.cell(5, i, h)
    style_header_row(ws, 5, 6)
    for idx, term in enumerate(terms):
        vals = [
            term["label"],
            weekday_desc(course["weekdays"]),
            term["dates"][0].isoformat(),
            term["dates"][-1].isoformat(),
            len(term["dates"]),
            term["color_name"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(6 + idx, c, v)
            cell.font = NORMAL_FONT
            cell.alignment = CENTER
            cell.border = THIN
            cell.fill = term["fill"]
    for i, w in enumerate([16, 28, 14, 14, 10, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    date_map = build_date_map(terms)
    for year, month in [(2026, 9), (2026, 10), (2026, 11), (2026, 12), (2027, 1)]:
        mws = wb.create_sheet(f"{year}年{month:02d}月")
        for col in range(1, 8):
            mws.column_dimensions[get_column_letter(col)].width = 12
        mws.column_dimensions["I"].width = 8
        for col in range(10, 16):
            mws.column_dimensions[get_column_letter(col)].width = 14
        write_calendar_month(mws, 1, year, month, date_map, course, terms)

    ws_all = wb.create_sheet("全部课节清单")
    ws_all["A1"] = f"{course['name']} · 全部课节清单"
    ws_all["A1"].font = TITLE_FONT
    for i, h in enumerate(["期次", "节次", "上课日期", "星期", "时间", "讲师", "颜色"], 1):
        ws_all.cell(3, i, h)
    style_header_row(ws_all, 3, 7)
    row_i = 4
    for term in terms:
        for i, d in enumerate(term["dates"], 1):
            vals = [
                term["label"],
                f"第{i}节",
                d.isoformat(),
                f"周{WEEKDAY_CN[d.weekday()]}",
                course["time"],
                course["teacher"],
                term["color_name"],
            ]
            for j, v in enumerate(vals, 1):
                cell = ws_all.cell(row_i, j, v)
                cell.font = NORMAL_FONT
                cell.alignment = CENTER
                cell.border = THIN
                cell.fill = term["fill"]
            row_i += 1
    for i, w in enumerate([16, 10, 14, 8, 14, 10, 8], 1):
        ws_all.column_dimensions[get_column_letter(i)].width = w

    return safe_save(wb, path)


# 各期剩余名额初始值（日常运营按前一日新增报名扣减后改此表或直接改xlsx）
DEFAULT_SEATS = {
    ("C01", 1): 2, ("C01", 2): 5, ("C01", 3): 8, ("C01", 4): 10, ("C01", 5): 12,
    ("C02", 1): 3, ("C02", 2): 6,
    ("C04", 1): 2, ("C04", 2): 4,
    ("C05", 1): 7, ("C05", 2): 9,
    ("C06", 1): 4, ("C06", 2): 5, ("C06", 3): 8,
}

SEATS_FILENAME = "各课程剩余名额.xlsx"


def write_remaining_seats_file(all_terms: dict[str, list[dict]]) -> Path:
    """生成统一剩余名额表；表底注明按前一天新增报名调整、每日一更。"""
    path = OUT_DIR / SEATS_FILENAME
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "剩余名额"

    ws["A1"] = "各课程实时剩余名额"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    today = date.today().isoformat()
    ws["A2"] = f"数据更新日期：{today}（每日一更）"
    ws["A2"].font = SUB_FONT
    ws.merge_cells("A2:F2")

    headers = ["课程编号", "课程名称", "期次", "开课日期", "结课日期", "剩余名额"]
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
    style_header_row(ws, 4, 6)

    row = 5
    for course in COURSES:
        cid = course["id"]
        for term in all_terms.get(cid, []):
            seats = DEFAULT_SEATS.get((cid, term["term_index"]), 0)
            vals = [
                cid,
                course["name"],
                term["label"],
                term["dates"][0].isoformat(),
                term["dates"][-1].isoformat(),
                seats,
            ]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row, c, v)
                cell.font = NORMAL_FONT
                cell.alignment = CENTER if c != 2 else LEFT
                cell.border = THIN
                if row % 2 == 1:
                    cell.fill = ALT_FILL
            row += 1

    note_row = row + 1
    ws.cell(note_row, 1, "说明：本表剩余名额需根据前一天新增的报名名额调整，每日一更。")
    ws.cell(note_row, 1).font = Font(name="微软雅黑", size=10, color="C00000")
    ws.cell(note_row, 1).alignment = LEFT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    ws.cell(note_row + 1, 1, "更新规则：每日根据前一日新增报名人数扣减对应期次剩余名额；满员记为0；勿手工编造未登记期次。")
    ws.cell(note_row + 1, 1).font = SUB_FONT
    ws.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=6)

    for i, w in enumerate([12, 22, 16, 14, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return safe_save(wb, path)


def delete_old_xlsx():
    for p in OUT_DIR.glob("*.xlsx"):
        # 保留人工每日更新的名额表，避免被课程重建脚本清空
        if p.name == SEATS_FILENAME:
            continue
        try:
            p.unlink()
            print("deleted:", p.name)
        except PermissionError:
            print("skip locked:", p.name)


def safe_save(wb, path: Path):
    try:
        wb.save(path)
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_新版" + path.suffix)
        wb.save(alt)
        print("原文件被占用，已另存为:", alt.name)
        return alt


def main():
    delete_old_xlsx()
    all_terms: dict[str, list[dict]] = {}
    for course in COURSES:
        terms = build_rolling_terms(course)
        all_terms[course["id"]] = terms
        print(course["id"], "terms=", len(terms), [t["label"] for t in terms])
        for t in terms:
            assert len(t["dates"]) == course["weeks"] * 2, (
                f"{course['id']} {t['label']} 课节数与周数不匹配: "
                f"{len(t['dates'])} != {course['weeks']*2}"
            )
            print(" ", t["label"], t["dates"][0], "->", t["dates"][-1], "n=", len(t["dates"]))
        p1 = write_content_file(course, terms)
        p2 = write_schedule_file(course, terms)
        print("created:", p1.name)
        print("created:", p2.name)

    # 若名额表尚不存在则生成；已存在则保留人工日报数据
    seats_path = OUT_DIR / SEATS_FILENAME
    if not seats_path.exists():
        p3 = write_remaining_seats_file(all_terms)
        print("created:", p3.name)
    else:
        print("keep existing:", SEATS_FILENAME)


if __name__ == "__main__":
    main()
